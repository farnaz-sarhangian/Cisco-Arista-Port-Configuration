import openpyxl

# Load the Excel file
workbook = openpyxl.load_workbook('input_excel_file.xlsx')
sheet = workbook.active

# Read the template text file
with open('template.txt', 'r') as template_file:
    template = template_file.read()

# Create a new text file for writing the output
with open('MOP.txt', 'w') as output_file:
    # Iterate through rows and generate output
    for row in sheet.iter_rows(min_row=2, values_only=True):
    
        PORTNUMBER, Desc, VLAN = row

        # Replace placeholders in the template with data
        output = template.replace('{PORT}', str(PORTNUMBER))
        output = output.replace('{DESC}', str(Desc))
        output = output.replace('{VLAN}', str(VLAN))

        # Write the modified template to the output file
        output_file.write(output + '\n')

        # Add an empty line after each entry
        output_file.write('\n') 

# Print a message when done
print("Output file has been saved as 'MOP.txt'")
